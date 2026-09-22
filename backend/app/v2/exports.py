import csv
import io
import json
import zipfile

from openpyxl import Workbook

from .contracts import HEADER_FIELDS, LINE_FIELDS, TAX_FIELDS
from .errors import V2Error


def safe_csv(value):
    if isinstance(value, str) and value.lstrip().startswith(
        ("=", "+", "-", "@", "\t", "\r")
    ):
        return "'" + value
    return value


def export(detail, format):
    if detail["status"] != "VERIFIED":
        raise V2Error("NOT_VERIFIED", "Only verified invoices can be exported.", 409)
    rid = detail["receipt_id"]
    header = {k: detail["fields"][k]["effective_value"] for k in HEADER_FIELDS}
    items = [
        dict(
            line_id=row["line_id"],
            **{k: row[k]["effective_value"] for k in LINE_FIELDS},
        )
        for row in detail["line_items"]
    ]
    taxes = [
        {k: row[k]["effective_value"] for k in TAX_FIELDS}
        for row in detail["tax_breakdown"]
    ]
    payload = dict(
        receipt_id=rid,
        review_status="VERIFIED",
        verified_at=detail["verified_at"],
        kie_run_id=detail["latest_kie_run_id"],
        header=header,
        line_items=items,
        tax_breakdown=taxes,
    )
    if format == "json":
        return (
            json.dumps(payload, ensure_ascii=False, allow_nan=False, indent=2).encode(),
            "application/json",
            f"{rid}.json",
        )
    tables = [
        (
            "Invoices",
            "invoice_headers.csv",
            ["receipt_id", "review_status", *HEADER_FIELDS],
            [[rid, "VERIFIED", *header.values()]],
        ),
        (
            "Line Items",
            "invoice_items.csv",
            ["receipt_id", "line_id", *LINE_FIELDS],
            [[rid, *row.values()] for row in items],
        ),
        (
            "Tax Groups",
            "invoice_taxes.csv",
            ["receipt_id", *TAX_FIELDS],
            [[rid, *row.values()] for row in taxes],
        ),
    ]
    output = io.BytesIO()
    if format == "csv":
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            for _, name, columns, rows in tables:
                text = io.StringIO(newline="")
                writer = csv.writer(text)
                writer.writerow(columns)
                writer.writerows([[safe_csv(v) for v in row] for row in rows])
                archive.writestr(name, text.getvalue().encode("utf-8-sig"))
        return output.getvalue(), "application/zip", f"{rid}-csv.zip"
    if format == "xlsx":
        wb = Workbook()
        wb.remove(wb.active)
        for title, _, columns, rows in tables:
            ws = wb.create_sheet(title)
            ws.append(columns)
            for row in rows:
                ws.append(row)
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                        cell.number_format = "@"
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        wb.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"{rid}.xlsx",
        )
    raise V2Error("INVALID_FORMAT", "Expected json, csv or xlsx.")
