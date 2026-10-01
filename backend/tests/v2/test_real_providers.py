"""Real PDF reader → KIE → persistence → review/export on synthetic invoices.

No model download, demo shim, OCR double or frozen evaluation data is used.
"""

import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from sqlalchemy import select

from backend.app.main import create_app
from backend.app.v2.errors import V2Error
from backend.app.v2.models import MachineRun
from backend.app.v2.processing import Processor
from backend.app.v2.providers import configured_kie, configured_reader
from test_workflow import kie, process


def invoice_pdf(pages=1):
    stream = io.BytesIO()
    document = canvas.Canvas(stream, pagesize=(600, 850))
    headers = [
        "Invoice template: 01GTKT0/001",
        "Invoice symbol: AA/26E",
        "Invoice number: 00001234",
        "Invoice date: 2026-09-24",
        "Seller name: Example Company",
        "Seller tax ID: 0101234567",
        "Seller address: 1 Sample Street",
        "Buyer name: Sample Buyer",
        "Buyer tax ID: 0201234567",
        "Currency: VND",
    ]
    xs = (35, 250, 330, 405, 490)
    items = [
        ("Office paper A4", "Ream", "2", "50000", "100000"),
        ("Ballpoint pen", "Piece", "3", "10000", "30000"),
    ]
    for page in range(pages):
        document.setFont("Helvetica", 10)
        for index, text in enumerate(headers):
            document.drawString(35, 810 - index * 22, text)
        for x, text in zip(
            xs, ("Description", "Unit", "Quantity", "Unit price", "Amount")
        ):
            document.drawString(x, 530, text)
        rows = items if pages == 1 else [items[page]]
        for index, row in enumerate(rows):
            for x, text in zip(xs, row):
                document.drawString(x, 505 - index * 25, text)
        for index, text in enumerate(
            ("Subtotal: 130000", "Total tax: 13000", "Grand total: 143000")
        ):
            document.drawString(35, 420 - index * 25, text)
        if page == pages - 1:
            for x, text in zip(
                (35, 250, 490), ("Tax rate", "Taxable amount", "Tax amount")
            ):
                document.drawString(x, 300, text)
            for x, text in zip((35, 250, 490), ("10%", "130000", "13000")):
                document.drawString(x, 275, text)
            document.drawString(35, 200, "Notes: synthetic demonstration invoice")
        document.showPage()
    document.save()
    return stream.getvalue()


@pytest.mark.parametrize("pages", [1, 2])
def test_real_pdf_to_verified_exports(svc, monkeypatch, pages):
    monkeypatch.delenv("V2_DOCUMENT_READER_CALLABLE", raising=False)
    monkeypatch.delenv("V2_KIE_CALLABLE", raising=False)
    client = TestClient(create_app(v2_service=svc))
    data = invoice_pdf(pages)
    response = client.post(
        "/api/v2/receipts", files={"file": ("synthetic.pdf", data, "application/pdf")}
    )
    assert response.status_code == 201
    rid = response.json()["receipt_id"]
    assert Processor(svc, configured_reader, configured_kie).process(rid)
    detail = client.get(f"/api/v2/receipts/{rid}").json()
    assert detail["status"] == "NEEDS_REVIEW"
    assert len(detail["fields"]) == 13
    assert detail["fields"]["invoice_number"]["effective_value"] == "00001234"
    assert detail["fields"]["seller_name"]["effective_value"] == "Example Company"
    assert detail["fields"]["seller_tax_id"]["effective_value"] == "0101234567"
    assert detail["fields"]["total_amount"]["effective_value"] == 143000
    assert [r["description"]["effective_value"] for r in detail["line_items"]] == [
        "Office paper A4",
        "Ballpoint pen",
    ]
    assert [r["amount"]["effective_value"] for r in detail["line_items"]] == [
        100000,
        30000,
    ]
    assert len(detail["tax_breakdown"]) == 1
    assert detail["tax_breakdown"][0]["rate"]["effective_value"] == "10%"
    assert client.get(f"/api/v2/receipts/{rid}/source").content == data
    evidence = client.get(f"/api/v2/receipts/{rid}/evidence").json()
    assert len(evidence["pages"]) == pages
    assert all(
        p["evidence"]["engine"]["name"] == "pdfium-text" for p in evidence["pages"]
    )
    assert client.get(f"/api/v2/receipts/{rid}/export?format=json").status_code == 409

    original_version = detail["version"]
    response = client.patch(
        f"/api/v2/receipts/{rid}/fields/invoice_number/correction",
        json={
            "expected_version": original_version,
            "value": "00000999",
            "status": "PRESENT",
        },
    )
    assert response.status_code == 200
    detail = response.json()
    assert detail["fields"]["invoice_number"]["normalized_value"] == "00001234"
    assert (
        client.post(
            f"/api/v2/receipts/{rid}/verify",
            json={"expected_version": original_version},
        ).status_code
        == 409
    )
    # Explicit human confirmation only on this test-created synthetic invoice.
    for section, rows in (
        ("fields", [("", detail["fields"])]),
        ("line-items", [(r["line_id"] + "/", r) for r in detail["line_items"]]),
        ("tax-groups", [(r["tax_id"] + "/", r) for r in detail["tax_breakdown"]]),
    ):
        for row_id, row in rows:
            for field, cell in row.items():
                if not isinstance(cell, dict) or not cell["effective_needs_review"]:
                    continue
                response = client.patch(
                    f"/api/v2/receipts/{rid}/{section}/{row_id}{field}/correction",
                    json={
                        "expected_version": detail["version"],
                        "value": cell["effective_value"],
                        "status": cell["effective_status"],
                    },
                )
                assert response.status_code == 200
                detail = response.json()
    response = client.post(
        f"/api/v2/receipts/{rid}/verify", json={"expected_version": detail["version"]}
    )
    assert response.status_code == 200 and response.json()["status"] == "VERIFIED"
    exported = client.get(f"/api/v2/receipts/{rid}/export?format=json").json()
    assert exported["header"]["invoice_number"] == "00000999"
    assert len(exported["line_items"]) == 2
    with zipfile.ZipFile(
        io.BytesIO(client.get(f"/api/v2/receipts/{rid}/export?format=csv").content)
    ) as archive:
        assert "00000999" in archive.read("invoice_headers.csv").decode("utf-8-sig")
        assert "Office paper A4" in archive.read("invoice_items.csv").decode(
            "utf-8-sig"
        )
    workbook = load_workbook(
        io.BytesIO(client.get(f"/api/v2/receipts/{rid}/export?format=xlsx").content)
    )
    assert workbook.sheetnames == ["Invoices", "Line Items", "Tax Groups"]
    assert "00000999" in [cell.value for row in workbook["Invoices"] for cell in row]
    with svc.sessions() as session:
        machine = session.scalar(
            select(MachineRun).where(MachineRun.run_id == detail["latest_kie_run_id"])
        )
        assert (
            machine.payload["fields"]["invoice_number"]["normalized_value"]
            == "00001234"
        )
    assert (
        client.get(f"/api/v2/receipts/{rid}/history").json()[-1]["kind"] == "VERIFIED"
    )


def test_reader_adapter_preserves_nonretryable_error():
    from uuid import uuid4
    from ai.document_reader.backend import read_document

    with pytest.raises(V2Error) as failure:
        read_document(
            b"%PDF-1.7 broken",
            content_type="application/pdf",
            receipt_id=str(uuid4()),
            ocr_run_id=str(uuid4()),
        )
    assert failure.value.code == "PDF_READ_FAILED"
    assert failure.value.status == 422 and not failure.value.retryable


@pytest.mark.parametrize(
    "corruption", ["raw_text", "row_sources", "row_review", "line_id"]
)
def test_backend_rejects_broken_cross_record_provenance(svc, corruption):
    def provider(evidence, **kwargs):
        result = kie(evidence, **kwargs)
        if corruption == "raw_text":
            result["fields"]["seller_name"]["raw_text"] = "invented source text"
        elif corruption == "row_sources":
            result["line_items"][0]["source_block_ids"] = ["foreign"]
        elif corruption == "row_review":
            result["line_items"][0]["machine_needs_review"] = True
        else:
            result["line_items"][0]["line_id"] = "page/1"
        return result

    detail = process(svc, provider=provider)
    assert detail["status"] == "FAILED"
    assert detail["processing_error"]["code"] == "SCHEMA_VALIDATION_FAILED"
    assert not detail["processing_error"]["retryable"]
    assert not detail["fields"]
