import copy
import io
import json
import zipfile
from concurrent.futures import ThreadPoolExecutor
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pypdf import PdfWriter
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker

from backend.app.main import create_app
from backend.app.storage.filesystem import FileSystemReceiptImageStorage
from backend.app.v2.contracts import HEADER_FIELDS, validate_result
from backend.app.v2.errors import V2Error
from backend.app.v2.models import Attempt, Base, MachineRun, Outbox
from backend.app.v2.processing import Processor
from backend.app.v2.service import InvoiceService


def image_bytes(fmt="PNG"):
    b = io.BytesIO()
    Image.new("RGB", (20, 20), "white").save(b, fmt)
    return b.getvalue()


def pdf_bytes():
    b = io.BytesIO()
    w = PdfWriter()
    w.add_blank_page(100, 100)
    w.write(b)
    return b.getvalue()


def value(v="00123"):
    return dict(
        raw_text=str(v),
        normalized_value=v,
        value_status="PRESENT",
        source_block_ids=["b1"],
        machine_needs_review=False,
    )


def reader(data, *, content_type, receipt_id, ocr_run_id):
    return dict(
        schema_version="1.3",
        receipt_id=str(receipt_id),
        ocr_run_id=str(ocr_run_id),
        engine=dict(name="test", version="1"),
        image=dict(width_px=20, height_px=20),
        blocks=[
            dict(
                block_id="b1",
                text="Công ty 00123",
                confidence=1,
                reading_order=0,
                polygon=[
                    dict(x=0, y=0),
                    dict(x=1, y=0),
                    dict(x=1, y=1),
                    dict(x=0, y=1),
                ],
            )
        ],
        average_confidence=1,
        duration_ms=1,
    )


def kie(evidence, *, kie_run_id):
    fields = {k: value("00123") for k in HEADER_FIELDS}
    fields.update(
        invoice_date=value("2026-09-21"),
        seller_name=value("Công ty Việt"),
        currency=value("VND"),
    )
    for k in ["subtotal", "tax_amount", "total_amount"]:
        fields[k] = value(100)
    return dict(
        schema_version="2.0",
        receipt_id=evidence["receipt_id"],
        kie_run_id=str(kie_run_id),
        source_ocr_run_id=evidence["ocr_run_id"],
        fields=fields,
        tax_breakdown=[
            dict(rate=value("10%"), taxable_amount=value(100), tax_amount=value(10))
        ],
        line_items=[
            dict(
                line_id="line-1",
                description=value("Cà phê"),
                unit=value("ly"),
                quantity=value(2),
                unit_price=value(50),
                amount=value(100),
            )
        ],
    )


@pytest.fixture
def svc(tmp_path):
    engine = create_engine(
        "sqlite:///" + str(tmp_path / "db.sqlite"), connect_args={"timeout": 20}
    )
    Base.metadata.create_all(engine)
    return InvoiceService(
        sessionmaker(engine, expire_on_commit=False),
        FileSystemReceiptImageStorage(tmp_path / "files"),
    )


def upload(svc):
    return svc.upload("invoice.png", "image/png", image_bytes())


def process(svc, r=None, provider=kie):
    r = r or upload(svc)
    Processor(svc, reader, provider).process(r["receipt_id"])
    return svc.detail(r["receipt_id"])


def test_upload_and_outbox(svc):
    r = upload(svc)
    assert r["status"] == "UPLOADED"
    assert svc.source(r["receipt_id"])[0] == image_bytes()
    with svc.sessions() as s:
        assert s.scalar(select(func.count()).select_from(Outbox)) == 1


@pytest.mark.parametrize(
    "name,mime,data",
    [
        ("x.jpg", "image/jpeg", image_bytes("JPEG")),
        ("x.png", "image/png", image_bytes()),
        ("x.pdf", "application/pdf", pdf_bytes()),
    ],
)
def test_valid_uploads(svc, name, mime, data):
    assert svc.upload(name, mime, data)["status"] == "UPLOADED"


@pytest.mark.parametrize(
    "name,mime,data",
    [
        ("x.png", "image/png", b""),
        ("x.png", "image/png", b"garbage"),
        ("x.pdf", "application/pdf", b"%PDF-fake"),
        ("x.jpg", "image/jpeg", image_bytes()),
        ("x.exe", "image/png", image_bytes()),
        ("x.png", "text/plain", image_bytes()),
    ],
)
def test_invalid_uploads(svc, name, mime, data):
    with pytest.raises(V2Error):
        svc.upload(name, mime, data)


def test_oversized_and_safe_filename(svc):
    svc.max_upload_bytes = 100
    with pytest.raises(V2Error):
        svc.upload("x.pdf", "application/pdf", b"x" * 101)
    svc.max_upload_bytes = 100000
    r = svc.upload("../../invoice.png", "image/png", image_bytes())
    assert r["original_filename"] == "invoice.png"


def test_pipeline_projects_all_v2_data(svc):
    r = process(svc)
    assert r["status"] == "NEEDS_REVIEW" and len(r["fields"]) == 13
    assert r["line_items"][0]["description"]["effective_value"] == "Cà phê"
    assert r["tax_breakdown"][0]["rate"]["effective_value"] == "10%"
    assert r["latest_ocr_run_id"] and r["latest_kie_run_id"]


def test_duplicate_delivery_and_claim_race(svc):
    r = upload(svc)
    with ThreadPoolExecutor(2) as pool:
        list(
            pool.map(
                lambda _: Processor(svc, reader, kie).process(r["receipt_id"]), range(2)
            )
        )
    Processor(svc, reader, kie).process(r["receipt_id"])
    with svc.sessions() as s:
        assert s.scalar(select(func.count()).select_from(Attempt)) == 1
        assert s.scalar(select(func.count()).select_from(MachineRun)) == 2


def test_failed_schema_retry_keeps_old_ocr_run(svc):
    def invalid(e, **kw):
        result = kie(e, **kw)
        result["fields"]["total_amount"]["normalized_value"] = "100"
        return result

    r = process(svc, provider=invalid)
    assert (
        r["status"] == "FAILED"
        and r["processing_error"]["code"] == "SCHEMA_VALIDATION_FAILED"
    )
    assert not r["processing_error"]["retryable"]
    with pytest.raises(V2Error):
        svc.retry(r["receipt_id"], r["version"])
    with svc.sessions() as s:
        assert s.scalar(select(func.count()).select_from(MachineRun)) == 1


def test_retry_new_runs_and_safe_error(svc):
    def broken(*a, **kw):
        raise RuntimeError("password=secret")

    r = process(svc, provider=broken)
    assert r["status"] == "FAILED" and "secret" not in json.dumps(r)
    old = r["latest_ocr_run_id"]
    svc.retry(r["receipt_id"], r["version"])
    r = process(svc, r)
    assert r["status"] == "NEEDS_REVIEW" and r["latest_ocr_run_id"] != old
    with svc.sessions() as s:
        assert s.scalar(select(func.count()).select_from(MachineRun)) == 3


def test_review_concurrency_and_immutable_machine(svc):
    r = process(svc)
    rid = r["receipt_id"]
    version = r["version"]
    r = svc.correct(rid, "header", "", "invoice_number", "000007", "PRESENT", version)
    assert r["fields"]["invoice_number"]["normalized_value"] == "00123"
    assert r["fields"]["invoice_number"]["effective_value"] == "000007"
    with pytest.raises(V2Error, match="version"):
        svc.correct(rid, "header", "", "invoice_number", "7", "PRESENT", version)
    r = svc.correct(rid, "line", "line-1", "quantity", 3, "PRESENT", r["version"])
    r = svc.correct(rid, "tax", "0", "rate", "8%", "PRESENT", r["version"])
    assert len(svc.history(rid)) == 3
    assert r["line_items"][0]["quantity"]["effective_value"] == 3
    assert r["tax_breakdown"][0]["rate"]["effective_value"] == "8%"


def test_verify_and_exports(svc):
    r = process(svc)
    rid = r["receipt_id"]
    with pytest.raises(V2Error):
        svc.export(rid, "json")
    r = svc.correct(
        rid, "header", "", "invoice_number", "000007", "PRESENT", r["version"]
    )
    r = svc.verify(rid, r["version"])
    assert r["status"] == "VERIFIED"
    with pytest.raises(V2Error):
        svc.verify(rid, r["version"])
    data, mime, name = svc.export(rid, "json")
    assert json.loads(data)["header"]["invoice_number"] == "000007"
    data, _, _ = svc.export(rid, "csv")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert set(z.namelist()) == {
            "invoice_headers.csv",
            "invoice_items.csv",
            "invoice_taxes.csv",
        }
        assert "000007" in z.read("invoice_headers.csv").decode("utf-8-sig")
        assert "Cà phê" in z.read("invoice_items.csv").decode("utf-8-sig")
    from openpyxl import load_workbook

    data, _, _ = svc.export(rid, "xlsx")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Invoices", "Line Items", "Tax Groups"]
    assert "000007" in [c.value for row in wb["Invoices"] for c in row]


def test_unknown_requires_explicit_review(svc):
    def unsure(e, **kw):
        result = kie(e, **kw)
        result["fields"]["buyer_name"].update(
            normalized_value=None, value_status="UNKNOWN"
        )
        return result

    r = process(svc, provider=unsure)
    with pytest.raises(V2Error):
        svc.verify(r["receipt_id"], r["version"])
    r = svc.correct(
        r["receipt_id"], "header", "", "buyer_name", None, "NOT_PRESENT", r["version"]
    )
    assert svc.verify(r["receipt_id"], r["version"])["status"] == "VERIFIED"


def test_nonpresent_nonnull_and_foreign_evidence_rejected():
    e = reader(b"", content_type="image/png", receipt_id=uuid4(), ocr_run_id=uuid4())
    result = kie(e, kie_run_id=uuid4())
    result["fields"]["buyer_name"]["value_status"] = "UNKNOWN"
    with pytest.raises(V2Error):
        validate_result(
            result, e, result["receipt_id"], result["kie_run_id"], e["ocr_run_id"]
        )
    result["fields"]["buyer_name"]["value_status"] = "PRESENT"
    result["line_items"][0]["amount"]["source_block_ids"] = ["foreign"]
    with pytest.raises(V2Error):
        validate_result(
            result, e, result["receipt_id"], result["kie_run_id"], e["ocr_run_id"]
        )


def test_api_workflow_and_v1_preserved(svc):
    client = TestClient(create_app(v2_service=svc))
    response = client.post(
        "/api/v2/receipts", files={"file": ("a.png", image_bytes(), "image/png")}
    )
    assert response.status_code == 201
    rid = response.json()["receipt_id"]
    assert client.get("/api/v2/receipts/" + rid).status_code == 200
    assert client.get("/api/v2/receipts/" + rid + "/source").content == image_bytes()
    assert client.get("/api/v2/receipts/not-a-uuid").status_code == 422
    assert client.get("/api/v1/receipts").status_code != 404


def test_recovery_fences_old_worker_and_retry(svc):
    from backend.app.v2.processing import recover_expired

    processor = Processor(svc, reader, kie)
    r = upload(svc)
    a = processor.claim(r["receipt_id"])
    with svc.sessions.begin() as s:
        s.get(Attempt, a["attempt_id"]).lease_until = 0
    assert recover_expired(svc) == 1
    with svc.sessions.begin() as s:
        assert not processor.fence(s, a, status="NEEDS_REVIEW")
    r = svc.detail(r["receipt_id"])
    assert r["status"] == "FAILED"
    svc.retry(r["receipt_id"], r["version"])
    r = process(svc, r)
    assert r["status"] == "NEEDS_REVIEW"
    with svc.sessions() as s:
        assert s.scalar(select(func.count()).select_from(Attempt)) == 2


def test_outbox_survives_broker_failure_and_republishes(svc):
    from backend.app.v2.processing import dispatch_pending

    r = upload(svc)

    def unavailable(rid):
        raise OSError("broker offline")

    assert dispatch_pending(svc, unavailable, now=100) == 0
    deliveries = []
    assert dispatch_pending(svc, deliveries.append, now=100) == 1
    assert deliveries == [r["receipt_id"]]
    assert dispatch_pending(svc, deliveries.append, now=120) == 0
    assert dispatch_pending(svc, deliveries.append, now=161) == 1
    process(svc, r)
    assert dispatch_pending(svc, deliveries.append, now=300) == 0


def test_exports_do_not_execute_formulas(svc):
    from openpyxl import load_workbook

    r = process(svc)
    r = svc.correct(
        r["receipt_id"], "header", "", "buyer_name", "=1+1", "PRESENT", r["version"]
    )
    svc.verify(r["receipt_id"], r["version"])
    data, _, _ = svc.export(r["receipt_id"], "xlsx")
    wb = load_workbook(io.BytesIO(data))
    found = [c for row in wb["Invoices"] for c in row if c.value == "=1+1"]
    assert found and found[0].data_type == "s"
    data, _, _ = svc.export(r["receipt_id"], "csv")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert "'=1+1" in z.read("invoice_headers.csv").decode("utf-8-sig")


def test_api_rejects_boolean_amount_and_exposes_evidence(svc):
    r = process(svc)
    client = TestClient(create_app(v2_service=svc))
    rid = r["receipt_id"]
    response = client.patch(
        f"/api/v2/receipts/{rid}/fields/total_amount/correction",
        json=dict(expected_version=r["version"], value=True, status="PRESENT"),
    )
    assert response.status_code == 422
    response = client.get(f"/api/v2/receipts/{rid}/evidence")
    assert (
        response.status_code == 200
        and response.json()["ocr_run_id"] == r["latest_ocr_run_id"]
    )


def test_missing_provider_is_explicit(svc, monkeypatch):
    from backend.app.v2.providers import configured_kie

    monkeypatch.delenv("V2_KIE_CALLABLE", raising=False)
    r = process(svc, provider=configured_kie)
    assert r["processing_error"]["code"] == "PROVIDER_NOT_CONFIGURED"


def test_multi_page_evidence_and_duplicate_line_rejected():
    from backend.app.v2.contracts import validate_evidence

    rid, ocr = str(uuid4()), str(uuid4())
    e = reader(b"", content_type="application/pdf", receipt_id=rid, ocr_run_id=ocr)
    envelope = dict(
        schema_version="document-2.0",
        receipt_id=rid,
        ocr_run_id=ocr,
        pages=[dict(page_index=0, evidence=e)],
    )
    validate_evidence(envelope, rid, ocr)
    result = kie(envelope, kie_run_id=uuid4())
    result["line_items"].append(copy.deepcopy(result["line_items"][0]))
    with pytest.raises(V2Error):
        validate_result(result, envelope, rid, result["kie_run_id"], ocr)


def test_openapi_matches_runtime():
    from pathlib import Path

    import yaml

    from scripts.export_openapi_v2 import document

    assert yaml.safe_load(Path("openapi/openapi-v2.yaml").read_text()) == document()


def test_review_line_ids_with_slashes(svc):
    def provider(e, **kw):
        result = kie(e, **kw)
        result["line_items"][0]["line_id"] = "page/1"
        return result

    r = process(svc, provider=provider)
    client = TestClient(create_app(v2_service=svc))
    response = client.patch(
        f"/api/v2/receipts/{r['receipt_id']}/line-items/page%2F1/description/correction",
        json={
            "expected_version": r["version"],
            "value": "Cà phê mới",
            "status": "PRESENT",
        },
    )
    assert response.status_code == 200
    assert (
        response.json()["line_items"][0]["description"]["effective_value"]
        == "Cà phê mới"
    )


@pytest.mark.parametrize("text", ["abc\x01def", "abc\ufffedef", "abc\ud800def"])
def test_reject_unexportable_correction(svc, text):
    r = process(svc)
    with pytest.raises(V2Error):
        svc.correct(
            r["receipt_id"], "header", "", "buyer_name", text, "PRESENT", r["version"]
        )


@pytest.mark.parametrize("kind", ["ocr", "kie"])
def test_non_finite_provider_output_fails_schema(svc, kind):
    def bad_reader(*a, **kw):
        result = reader(*a, **kw)
        result["blocks"][0]["confidence"] = float("nan")
        return result

    def bad_kie(*a, **kw):
        result = kie(*a, **kw)
        result["fields"]["buyer_name"]["confidence"] = float("nan")
        return result

    r = upload(svc)
    Processor(
        svc, bad_reader if kind == "ocr" else reader, bad_kie if kind == "kie" else kie
    ).process(r["receipt_id"])
    r = svc.detail(r["receipt_id"])
    assert (
        r["status"] == "FAILED"
        and r["processing_error"]["code"] == "SCHEMA_VALIDATION_FAILED"
    )
